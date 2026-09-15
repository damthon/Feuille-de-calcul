"""Export PDF et Excel de la note de calcul — extra ``[export]`` (openpyxl, reportlab).

Réutilise directement ``ResultatCalcul`` (``calcul.py``) — ni ce module
ni ``rapport.py`` ne recalculent quoi que ce soit ; ils mettent
seulement en forme des résultats déjà obtenus, dans des formats
différents (Markdown pour ``rapport.py``, tableur pour
``exporter_excel``, document imprimable pour ``exporter_pdf``).

PDF via ``reportlab`` plutôt que ``weasyprint`` (les deux étaient
envisagés au plan de conception, §10) : bibliothèque pure Python (à
quelques extensions C près), sans dépendance système native
(Cairo/Pango/GTK) — plus simple à empaqueter de façon fiable avec
PyInstaller (lot 11), cohérent avec l'objectif d'un exécutable léger
(§3 du plan).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from mur_contre_terre.calcul import ResultatCalcul
from mur_contre_terre.donnees.projet import Projet
from mur_contre_terre.donnees.sol import TypePoussee
from mur_contre_terre.unites import en_deg, en_kN, en_kN_m2, en_kN_m3, en_MPa

_AVERTISSEMENT_TRANCHANT = (
    "Résistance à l'effort tranchant (SIA 262 §4.3.3) : formule harmonisée EN 1992-1-1 §6.2.2, "
    "faute d'extrait SIA 262 vérifié pour ce lot — à confirmer avant toute justification de projet."
)


def _lignes_verifications(resultat: ResultatCalcul) -> list[list[str | float]]:
    lignes: list[list[str | float]] = [
        ["z [m]", "h [cm]", "As terre [cm²/m]", "As intérieur [cm²/m]", "V_Ed [kN]", "V_Rd [kN]", "Effort tranchant"]
    ]
    for v in resultat.verifications:
        lignes.append(
            [
                round(v.position, 2),
                round(v.epaisseur * 100, 1),
                round(v.armature_terre * 1e4, 2),
                round(v.armature_interieur * 1e4, 2),
                round(en_kN(v.effort_tranchant_ed), 1),
                round(en_kN(v.effort_tranchant_rd), 1),
                "OK" if v.verdict_tranchant else "Insuffisant",
            ]
        )
    return lignes


# --------------------------------------------------------------------- Excel


def _feuille_donnees(classeur: Workbook, projet: Projet) -> Worksheet:
    feuille = classeur.active
    feuille.title = "Données"
    g, s, m = projet.geometrie, projet.sol, projet.materiaux
    type_poussee = "active" if s.type_poussee is TypePoussee.ACTIF else "au repos"
    lignes = [
        ("Projet", projet.nom),
        ("Hauteur H [m]", g.hauteur),
        ("Épaisseur en pied [m]", g.ep_base),
        ("Épaisseur en tête [m]", g.ep_couronnement),
        ("Débord de semelle [m]", g.debord_semelle),
        ("Épaisseur de semelle [m]", g.ep_semelle),
        ("Poids volumique γ [kN/m³]", round(en_kN_m3(s.gamma), 2)),
        ("Angle de frottement φ′ [°]", round(en_deg(s.phi), 2)),
        ("Cohésion c′ [kN/m²]", round(en_kN_m2(s.c), 2)),
        ("Type de poussée", type_poussee),
        ("Classe de béton", f"fck = {en_MPa(m.beton.fck):.0f} MPa"),
        ("Enrobage côté terre [mm]", m.enrobage_terre * 1000.0),
        ("Enrobage côté intérieur [mm]", m.enrobage_interieur * 1000.0),
    ]
    for ligne, (cle, valeur) in enumerate(lignes, start=1):
        feuille.cell(row=ligne, column=1, value=cle).font = Font(bold=True)
        feuille.cell(row=ligne, column=2, value=valeur)
    feuille.column_dimensions["A"].width = 28
    feuille.column_dimensions["B"].width = 24
    return feuille


def _feuille_verifications(classeur: Workbook, resultat: ResultatCalcul) -> Worksheet:
    feuille = classeur.create_sheet("Vérification de section")
    for ligne in _lignes_verifications(resultat):
        feuille.append(ligne)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for colonne in feuille.columns:
        lettre = colonne[0].column_letter
        feuille.column_dimensions[lettre].width = 16
    feuille.append([])
    feuille.append([_AVERTISSEMENT_TRANCHANT])
    return feuille


def _feuille_combinaisons(classeur: Workbook, resultat: ResultatCalcul) -> Worksheet:
    feuille = classeur.create_sheet("Combinaisons ELU")
    feuille.append(["Combinaison", "Facteurs"])
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for combinaison in resultat.combinaisons_elu:
        facteurs = ", ".join(f"{nom} × {f:.2f}" for nom, f in combinaison.facteurs.items())
        feuille.append([combinaison.nom, facteurs])
    feuille.column_dimensions["A"].width = 30
    feuille.column_dimensions["B"].width = 60
    return feuille


def exporter_excel(projet: Projet, resultat: ResultatCalcul, chemin: str | Path) -> None:
    """Classeur Excel : données du projet, combinaisons ELU, vérification de section (une feuille chacune)."""
    classeur = Workbook()
    _feuille_donnees(classeur, projet)
    _feuille_combinaisons(classeur, resultat)
    _feuille_verifications(classeur, resultat)
    classeur.save(str(chemin))


# ----------------------------------------------------------------------- PDF


def _table_pdf(lignes: list[list[str]], entete: bool = True) -> Table:
    table = Table(lignes, hAlign="LEFT")
    style = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D3C6")),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if entete:
        style += [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ECE8DF")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ]
    table.setStyle(TableStyle(style))
    return table


def exporter_pdf(projet: Projet, resultat: ResultatCalcul, chemin: str | Path) -> None:
    """Note de calcul PDF : géométrie, sol, matériaux, combinaisons ELU, vérification de section."""
    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(
        str(chemin), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm
    )
    elements: list = [Paragraph(f"Note de calcul — {projet.nom}", styles["Title"]), Spacer(1, 0.5 * cm)]

    g, s, m = projet.geometrie, projet.sol, projet.materiaux
    type_poussee = "active" if s.type_poussee is TypePoussee.ACTIF else "au repos"
    elements.append(Paragraph("Géométrie et sol", styles["Heading2"]))
    elements.append(
        _table_pdf(
            [
                ["Grandeur", "Valeur"],
                ["Hauteur H", f"{g.hauteur:.2f} m"],
                ["Épaisseur en pied / en tête", f"{g.ep_base:.3f} / {g.ep_couronnement:.3f} m"],
                ["Débord et épaisseur de semelle", f"{g.debord_semelle:.2f} / {g.ep_semelle:.2f} m"],
                ["Poids volumique γ", f"{en_kN_m3(s.gamma):.1f} kN/m³"],
                ["Angle de frottement φ′", f"{en_deg(s.phi):.1f} °"],
                ["Type de poussée", f"{type_poussee} (SIA 261 §4.3.2)"],
                ["Béton / enrobages", f"fck = {en_MPa(m.beton.fck):.0f} MPa, {m.enrobage_terre*1000:.0f}/{m.enrobage_interieur*1000:.0f} mm"],
            ]
        )
    )
    elements.append(Spacer(1, 0.5 * cm))

    elements.append(Paragraph("Combinaisons ELU (SIA 260 §4.4.3.2)", styles["Heading2"]))
    lignes_combi = [["Combinaison", "Facteurs"]] + [
        [c.nom, ", ".join(f"{nom} × {f:.2f}" for nom, f in c.facteurs.items())] for c in resultat.combinaisons_elu
    ]
    elements.append(_table_pdf(lignes_combi))
    elements.append(Spacer(1, 0.5 * cm))

    elements.append(Paragraph("Vérification de section (SIA 262 §4.1, §4.3.3)", styles["Heading2"]))
    lignes = [[str(v) for v in ligne] for ligne in _lignes_verifications(resultat)]
    elements.append(_table_pdf(lignes))
    elements.append(Spacer(1, 0.3 * cm))
    elements.append(Paragraph(_AVERTISSEMENT_TRANCHANT, styles["Italic"]))

    document.build(elements)
