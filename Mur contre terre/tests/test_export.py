import openpyxl
import pytest

from mur_contre_terre.calcul import calculer
from mur_contre_terre.donnees.appuis import ConditionsAppui, TypeAppuiPied, TypeAppuiTete
from mur_contre_terre.donnees.charges import CasDeCharge, CategorieAction, TypeCharge
from mur_contre_terre.donnees.geometrie import Geometrie
from mur_contre_terre.donnees.materiaux import Acier, Beton, Materiaux
from mur_contre_terre.donnees.projet import Projet
from mur_contre_terre.donnees.sol import Sol, TypePoussee
from mur_contre_terre.export import exporter_excel, exporter_pdf
from mur_contre_terre.unites import deg, kN, kN_m3


def _projet() -> Projet:
    return Projet(
        nom="Test export",
        geometrie=Geometrie(hauteur=3.0, ep_base=0.30, ep_couronnement=0.20, debord_semelle=0.8, ep_semelle=0.4),
        sol=Sol(gamma=kN_m3(18), phi=deg(30), delta=deg(20), type_poussee=TypePoussee.ACTIF),
        appuis=ConditionsAppui(pied=TypeAppuiPied.ENCASTREMENT, tete=TypeAppuiTete.LIBRE),
        materiaux=Materiaux(
            beton=Beton.depuis_classe("C30/37"), acier=Acier.b500b(), enrobage_terre=0.05, enrobage_interieur=0.04
        ),
        charges=(
            CasDeCharge(nom="Poids propre", type=TypeCharge.POIDS_PROPRE, categorie=CategorieAction.G, valeur=0.0),
            CasDeCharge(nom="Poussee", type=TypeCharge.POUSSEE_TERRES, categorie=CategorieAction.G, valeur=0.0),
            CasDeCharge(
                nom="Charge tete", type=TypeCharge.CHARGE_TETE, categorie=CategorieAction.Q, valeur=kN(30),
                psi0=0.7, psi1=0.5, psi2=0.3,
            ),
        ),
        finesse_maillage=0.5,
    )


def test_exporter_excel_cree_les_trois_feuilles(tmp_path):
    projet = _projet()
    resultat = calculer(projet)
    chemin = tmp_path / "note.xlsx"

    exporter_excel(projet, resultat, chemin)

    assert chemin.exists()
    classeur = openpyxl.load_workbook(chemin)
    assert classeur.sheetnames == ["Données", "Combinaisons ELU", "Vérification de section"]


def test_exporter_excel_donnees_contient_le_nom_du_projet(tmp_path):
    projet = _projet()
    resultat = calculer(projet)
    chemin = tmp_path / "note.xlsx"

    exporter_excel(projet, resultat, chemin)

    classeur = openpyxl.load_workbook(chemin)
    valeurs = [cellule.value for ligne in classeur["Données"].iter_rows() for cellule in ligne]
    assert "Test export" in valeurs


def test_exporter_excel_verifications_a_une_ligne_par_element(tmp_path):
    projet = _projet()
    resultat = calculer(projet)
    chemin = tmp_path / "note.xlsx"

    exporter_excel(projet, resultat, chemin)

    feuille = openpyxl.load_workbook(chemin)["Vérification de section"]
    # 1 ligne d'en-tête + une ligne par élément + une ligne vide + une ligne d'avertissement
    lignes_non_vides = [ligne for ligne in feuille.iter_rows(values_only=True) if any(v is not None for v in ligne)]
    assert len(lignes_non_vides) == 1 + len(resultat.verifications) + 1


def test_exporter_pdf_cree_un_fichier_pdf_valide(tmp_path):
    projet = _projet()
    resultat = calculer(projet)
    chemin = tmp_path / "note.pdf"

    exporter_pdf(projet, resultat, chemin)

    assert chemin.exists()
    assert chemin.read_bytes()[:5] == b"%PDF-"
    assert chemin.stat().st_size > 500
